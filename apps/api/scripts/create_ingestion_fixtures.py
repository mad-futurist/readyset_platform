"""Generate the committed, entirely synthetic M2 parser corpus."""
from pathlib import Path

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

root = Path(__file__).resolve().parents[1] / "tests" / "fixtures" / "ingestion"
root.mkdir(parents=True, exist_ok=True)
writer = PdfWriter()
for text in ["Security policy|Rotate keys every 30 days.", "", "Support policy|Contact the support desk on page three."]:
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(("BT /F1 12 Tf 50 740 Td " + " 0 -20 Td ".join(f"({line}) Tj" for line in text.split("|")) + " ET").encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(stream)
with (root / "policies.pdf").open("wb") as target:
    writer.write(target)
document = Document()
document.add_heading("Deployment", level=1)
document.add_paragraph("Deploy after reviewing the release checklist.")
document.add_heading("Backups", level=2)
table = document.add_table(rows=2, cols=2)
table.cell(0, 0).text, table.cell(0, 1).text = "Region", "Retention"
table.cell(1, 0).text, table.cell(1, 1).text = "Paris", "30 days"
document.add_paragraph("Unicode: café, Україна, 東京.")
document.save(str(root / "deployment.docx"))
(root / "deployment.md").write_text("# Deployment\n\nReview the release checklist.\n\n## Backups\n\nKeep backups for 30 days.\n\n```sh\necho 'synthetic example'\n```\n\n## Empty section\n", encoding="utf-8")
(root / "unicode.txt").write_text("Hello café.\nУкраїна and 東京.\n\nSecond paragraph.\n", encoding="utf-8")
(root / "malformed.pdf").write_bytes(b"%PDF-1.7\nThis is not a PDF object graph.\n%%EOF\n")
(root / "malformed.docx").write_bytes(b"PK\x03\x04not a valid zip")
